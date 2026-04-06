# -*- coding: utf-8 -*-
import os
import tkinter as tk
import openpyxl
from openpyxl.worksheet.protection import SheetProtection
from tkinter import messagebox
from config import SETTING_DIR, SHEET_PASS

# Define the correct password here (keep it consistent)
CORRECT_PASSWORD = "123"
sheet_pass = SHEET_PASS  # Use the password from config.py
MAIN_PATH = SETTING_DIR
MAIN_FILE = os.path.join(MAIN_PATH, 'Daily Template.xlsx')
SLUDGE_FILE = os.path.join(MAIN_PATH, 'Dry Sludge Template.xlsx')
EXTERNAL_FILE = os.path.join(MAIN_PATH, 'External Report Template.xlsx')


def load_existing_settings():
    """Load current settings from Excel file"""
    try:
        wb = openpyxl.load_workbook(MAIN_FILE, data_only=True)
        sheet_main = wb['Main']
        wb1 = openpyxl.load_workbook(SLUDGE_FILE, data_only=True)
        wb1_sheet_main = wb1['Main']
        wb2 = openpyxl.load_workbook(EXTERNAL_FILE, data_only=True)
        wb2_sheet_main = wb2['Main']
        settings = {
            'company_governorate_en': 'Faiyum',
            'company_governorate_ar': 'الفيوم',
            'plant_name_en': 'Tamiya',
            'plant_name_ar': 'طامية',
            'technical_manager': 'Mohamad',
            'design_flow': '10,000',
            'actual_flow': '12,000'
        }

        # Try to extract actual values if they exist
        try:
            cell_i4 = str(sheet_main['I4'].value or '')
            cell_i2 = str(wb1_sheet_main['I2'].value or '')
            cell_i5 = str(wb2_sheet_main['I2'].value or '')
            prefix = 'شركة مياه الشرب والصرف الصحى ب'
            if cell_i4.startswith(prefix):
                settings['company_governorate_ar'] = cell_i4[len(
                    prefix):].strip()
            else:
                settings['company_governorate_ar'] = cell_i4.strip()
            if cell_i2.startswith(prefix):
                settings['company_governorate_ar'] = cell_i2[len(
                    prefix):].strip()
            else:
                settings['company_governorate_ar'] = cell_i2.strip()
            if cell_i5.startswith(prefix):
                settings['company_governorate_ar'] = cell_i5[len(
                    prefix):].strip()
            else:
                settings['company_governorate_ar'] = cell_i5.strip()
        except:
            pass

        try:
            cell_b4 = str(sheet_main['B4'].value or '')
            cell_d2 = str(wb1_sheet_main['D2'].value or '')
            cell_a2 = str(wb2_sheet_main['A2'].value or '')

            suffix = ' Drinking Water and Sanitation Company'
            if cell_b4.endswith(suffix):
                settings['company_governorate_en'] = cell_b4[:-
                                                             len(suffix)].strip().title()
            else:
                settings['company_governorate_en'] = cell_b4.strip().title()

            if cell_d2.endswith(suffix):
                settings['company_governorate_en'] = cell_d2[:-
                                                             len(suffix)].strip().title()
            else:
                settings['company_governorate_en'] = cell_d2.strip().title()

            if cell_a2.endswith(suffix):
                settings['company_governorate_en'] = cell_a2[:-
                                                             len(suffix)].strip().title()
            else:
                settings['company_governorate_en'] = cell_a2.strip().title()

        except:
            pass

        try:
            cell_j7 = str(sheet_main['J7'].value or '')
            cell_i5 = str(wb1_sheet_main['I5'].value or '')
            prefix = 'معمل محطة معالجة '
            if cell_j7.startswith(prefix):
                settings['plant_name_ar'] = cell_j7[len(prefix):].strip()
            else:
                settings['plant_name_ar'] = cell_j7.strip()
            if cell_i5.startswith(prefix):
                settings['plant_name_ar'] = cell_i5[len(prefix):].strip()
            else:
                settings['plant_name_ar'] = cell_i5.strip()
        except:
            pass

        try:
            cell_a7 = str(sheet_main['A7'].value or '')
            cell_a6 = str(wb1_sheet_main['A6'].value or '')

            suffix = ' Sanitation Plant laboratory'
            if cell_a7.endswith(suffix):
                settings['plant_name_en'] = cell_a7[:-
                                                    len(suffix)].strip().title()
            else:
                settings['plant_name_en'] = cell_a7.strip().title()
            if cell_a6.endswith(suffix):
                settings['plant_name_en'] = cell_a6[:-
                                                    len(suffix)].strip().title()
            else:
                settings['plant_name_en'] = cell_a6.strip().title()
        except:
            pass

        try:
            sheet_cod = wb['COD']
            cell_i8 = str(sheet_cod['I8'].value or '')
            cell_e19 = str(wb1_sheet_main['E19'].value or '')
            cell_a6 = str(wb2_sheet_main['COD']['I8'].value or '')
            prefix = 'ANALYST :    '
            if cell_i8.startswith(prefix):
                settings['technical_manager'] = cell_i8[len(
                    prefix):].strip().title()
            else:
                settings['technical_manager'] = cell_i8.strip().title()
            if cell_e19.startswith(prefix):
                settings['technical_manager'] = cell_e19[len(
                    prefix):].strip().title()
            else:
                settings['technical_manager'] = cell_e19.strip().title()

            if cell_a6.startswith(prefix):
                settings['technical_manager'] = cell_a6[len(
                    prefix):].strip().title()
            else:
                settings['technical_manager'] = cell_a6.strip().title()

        except:
            pass

        try:
            sheet_bod = wb['BOD5']
            settings['actual_flow'] = str(sheet_bod['C30'].value or '12,000')
        except:
            pass

        wb.close()
        return settings
    except Exception as e:
        print(f"Failed to load existing settings: {e}")
        return None


def update_excel_cells(settings_dict):
    """Centralized function to update Excel cells"""
    try:
        wb = openpyxl.load_workbook(MAIN_FILE)
        wb1 = openpyxl.load_workbook(SLUDGE_FILE)

        # Update main sheet
        sheet_main = wb['Main']
        wb1_sheet_main = wb1['Main']
        # unprotect
        sheet_main.protection = SheetProtection(sheet=False)
        wb1_sheet_main.protection = SheetProtection(sheet=False)

        sheet_main['I4'] = f"شركة مياه الشرب والصرف الصحى ب{settings_dict['company_governorate_ar']}"
        sheet_main['A4'] = f"{settings_dict['company_governorate_en']} Drinking Water and Sanitation Company"
        sheet_main['J7'] = f"محطة معالجة {settings_dict['plant_name_ar']}"
        wb1_sheet_main['i7'] = f"محطة معالجة {settings_dict['plant_name_ar']}"
        sheet_main['A7'] = f"{settings_dict['plant_name_en']} Treatment Plant"
        wb1_sheet_main['b7'] = f"{settings_dict['plant_name_en']} Treatment Plant"
        # re-protect with password
        sheet_main.protection = SheetProtection(
            sheet=True, password=sheet_pass)
        wb1_sheet_main.protection = SheetProtection(
            sheet=True, password=sheet_pass)

        # Update COD sheet
        sheet_cod = wb['COD']
        sheet_cod.protection = SheetProtection(sheet=False)
        sheet_cod['I8'] = f"ANALYST :    {settings_dict['technical_manager']}"
        sheet_cod.protection = SheetProtection(sheet=True, password=sheet_pass)

        # Update BOD5 sheet
        sheet_bod = wb['BOD5']
        sheet_bod.protection = SheetProtection(sheet=False)
        sheet_bod['C30'] = settings_dict['actual_flow']
        sheet_bod.protection = SheetProtection(sheet=True, password=sheet_pass)

        wb.save(MAIN_FILE)
        wb.close()
        wb1.save(SLUDGE_FILE)
        wb1.close()
        return True
    except PermissionError:
        raise Exception("Please close the Excel file before saving settings.")
    except FileNotFoundError:
        raise Exception(f"Excel file not found at: {MAIN_FILE}")
    except Exception as e:
        raise Exception(f"Excel update failed: {e}")


def validate_inputs(entries):
    """Validate all input fields"""
    errors = []

    # Required fields validation
    required_fields = ['Co', 'Co_AR', 'Plant', 'Plant_Ar', 'Technical_Manager']
    for field in required_fields:
        if not entries[field].get().strip():
            errors.append(f"{field} cannot be empty")

    # Numeric fields validation
    numeric_fields = ['Plant_Designing_Flow', 'Plant_Actual_Flow']
    for field in numeric_fields:
        value = entries[field].get().replace(',', '')
        if not value.replace('.', '').isdigit():
            errors.append(f"{field} must be numeric")

    return errors


def create_input_section(parent, label_text, default_value="", arabic_font=None):
    """Create standardized input sections"""
    frame = tk.Frame(parent)
    frame.pack(pady=8, fill='x', padx=20)
    # Ensure a valid font tuple is provided to satisfy type-checkers
    safe_font = arabic_font if arabic_font is not None else ("Arial", 12)

    lbl = tk.Label(frame, text=label_text, font=safe_font, anchor='w')
    lbl.pack(side='top', anchor='w')

    entry = tk.Entry(frame, width=60, font=("Arial", 11))
    entry.pack(side='top', fill='x', pady=3)
    entry.insert(0, default_value)

    return entry


def open_settings_form(root_window, arabic_font, pwd_window):
    """Opens the actual settings form after password verification."""
    pwd_window.destroy()

    settings_window = tk.Toplevel(root_window)
    settings_window.title("إعدادات التطبيق")
    settings_window.geometry("800x650")

    # Set minimum size
    settings_window.minsize(600, 700)

    # Create main container with scrollbar
    main_frame = tk.Frame(settings_window)
    main_frame.pack(fill='both', expand=True, padx=5, pady=5)

    canvas = tk.Canvas(main_frame, highlightthickness=0)
    scrollbar = tk.Scrollbar(
        main_frame, orient="vertical", command=canvas.yview)
    scrollable_frame = tk.Frame(canvas)

    scrollable_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    # Add title
    title_label = tk.Label(
        scrollable_frame,
        text="تعديل إعدادات التقرير اليومي",
        font=(arabic_font[0], 14, "bold"),
        fg="darkblue"
    )
    title_label.pack(pady=(10, 20))

    # Load existing settings
    current_settings = load_existing_settings() or {
        'company_governorate_en': 'Faiyum',
        'company_governorate_ar': 'الفيوم',
        'plant_name_en': 'Tamiya',
        'plant_name_ar': 'طامية',
        'technical_manager': 'Mohamad',
        'design_flow': '10,000',
        'actual_flow': '12,000'
    }

    # Create input fields
    entries = {}
    fields = [
        ('Co', 'Type the Company Governorate Name in English:',
         current_settings['company_governorate_en']),
        ('Co_AR', 'Type the Company Governorate Name in Arabic:',
         current_settings['company_governorate_ar']),
        ('Plant', 'Type the Plant Name in English:',
         current_settings['plant_name_en']),
        ('Plant_Ar', 'Type the Plant Name in Arabic:',
         current_settings['plant_name_ar']),
        ('Technical_Manager', 'Type the Technical Manager Name:',
         current_settings['technical_manager']),
        ('Plant_Designing_Flow', 'Enter Plant Designing Flow (m³/day):',
         current_settings['design_flow']),
        ('Plant_Actual_Flow', 'Enter Plant Actual Flow (m³/day):',
         current_settings['actual_flow'])
    ]

    for field_name, label_text, default_value in fields:
        entry = create_input_section(
            scrollable_frame, label_text, default_value, arabic_font)
        entries[field_name] = entry

    # Create list of entries in order for proper navigation
    entries_list: list = [entries[field[0]] for field in fields]

    # Focus on the first entry field after a tiny delay (more reliable)
    settings_window.after(100, lambda: entries_list[0].focus_set())

    # Bind Enter keys for navigation
    for i in range(len(entries_list) - 1):
        current_entry = entries_list[i]
        next_entry = entries_list[i + 1]
        current_entry.bind('<Return>', lambda e,
                           nxt=next_entry: nxt.focus_set())

    # Last entry: Bind to the save function
    entries_list[-1].bind('<Return>', lambda e: save_settings())

    def save_settings():
        # Validate inputs
        errors = validate_inputs(entries)
        if errors:
            messagebox.showerror("خطأ في التحقق", "\n".join(errors))
            return

        # Prepare settings data
        settings_data = {
            'company_governorate_en': entries['Co'].get().strip().title(),
            'company_governorate_ar': entries['Co_AR'].get().strip(),
            'plant_name_en': entries['Plant'].get().strip().title(),
            'plant_name_ar': entries['Plant_Ar'].get().strip(),
            'technical_manager': entries['Technical_Manager'].get().strip().title(),
            'design_flow': '10,000',  # Default value
            'actual_flow': entries['Plant_Actual_Flow'].get().strip()
        }

        # Ask for confirmation
        if not messagebox.askyesno("تأكيد", "هل تريد حفظ التغييرات؟"):
            return

        try:
            # Update Excel file
            update_excel_cells(settings_data)
            messagebox.showinfo("نجاح", "تم حفظ الإعدادات بنجاح.")
            root_window.deiconify()  # type: ignore
            settings_window.destroy()
        except Exception as e:
            messagebox.showerror("خطأ", f"فشل في حفظ الإعدادات: {str(e)}")

    def cancel_settings():
        settings_window.destroy()
        root_window.deiconify()  # type: ignore

    # Button frame
    button_frame = tk.Frame(settings_window)
    button_frame.pack(fill='x', pady=10)

    # Buttons
    btn_save = tk.Button(
        button_frame,
        text="حفظ الإعدادات",
        command=save_settings,
        font=arabic_font,
        bg='green',
        fg='white',
        padx=15,
        pady=8,
        relief='raised',
        bd=2
    )
    btn_save.pack(side='right', padx=10)

    btn_cancel = tk.Button(
        button_frame,
        text="إلغاء",
        command=cancel_settings,
        font=arabic_font,
        bg='red',
        fg='white',
        padx=10,
        pady=8,
        relief='raised',
        bd=2
    )
    btn_cancel.pack(side='left', padx=10)

    # back to main menu button
    btn_back = tk.Button(
        button_frame,
        text="العودة للقائمة الرئيسية",
        command=lambda: [settings_window.destroy(), root_window.deiconify()],
        font=arabic_font,
        bg='orange',
        fg='white',
        padx=20,
        pady=8,
        relief='raised',
        bd=2
    )
    btn_back.pack(side='left', padx=10)

    # back to previous page button
    btn_previous = tk.Button(
        button_frame,
        text="الرجوع إلى الصفحة السابقة",
        command=lambda: [settings_window.destroy(
        ), prompt_for_password(root_window, arabic_font)],
        font=arabic_font,
        bg='orange',
        fg='white',
        padx=30,
        pady=8,
        relief='raised',
        bd=2
    )
    btn_previous.pack(side='left', padx=10)

    # Pack canvas and scrollbar
    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")


def prompt_for_password(root_window, arabic_font):
    """Displays the password prompt window."""
    pwd_window = tk.Toplevel(root_window)
    pwd_window.title("كلمة المرور")
    pwd_window.geometry("350x200")
    pwd_window.grab_set()
    root_window.withdraw()

    # Center window
    pwd_window.update_idletasks()
    width = pwd_window.winfo_width()
    height = pwd_window.winfo_height()
    x = (pwd_window.winfo_screenwidth() // 2) - (width // 2)
    y = (pwd_window.winfo_screenheight() // 2) - (height // 2)
    pwd_window.geometry(f'{width}x{height}+{x}+{y}')

    def verify_password(event=None):
        entered_password = entry_password.get()
        if entered_password == CORRECT_PASSWORD:
            open_settings_form(root_window, arabic_font, pwd_window)
        else:
            messagebox.showerror("خطأ", "كلمة المرور غير صحيحة.")
            entry_password.delete(0, tk.END)
            entry_password.focus_set()

    def cancel_password():
        pwd_window.destroy()
        root_window.deiconify()

    # Title
    lbl_title = tk.Label(
        pwd_window,
        text="وصول الإعدادات",
        font=(arabic_font[0], 14, "bold"),
        fg="darkred"
    )
    lbl_title.pack(pady=(15, 10))

    lbl_password = tk.Label(
        pwd_window,
        text="أدخل كلمة المرور:",
        font=arabic_font
    )
    lbl_password.pack(pady=5)

    entry_password = tk.Entry(
        pwd_window,
        show="●",
        font=("Arial", 12),
        width=25,
        justify='center'
    )
    entry_password.pack(pady=10, ipady=5)
    entry_password.focus_set()

    entry_password.bind('<Return>', verify_password)

    btn_frame = tk.Frame(pwd_window)
    btn_frame.pack(pady=10)

    btn_verify = tk.Button(
        btn_frame,
        text="دخول",
        command=verify_password,
        font=arabic_font,
        bg='blue',
        fg='white',
        padx=25,
        pady=5,
        relief='raised',
        bd=2
    )
    btn_verify.pack(side='left', padx=5)

    btn_cancel = tk.Button(
        btn_frame,
        text="إلغاء",
        command=cancel_password,
        font=arabic_font,
        bg='gray',
        fg='white',
        padx=25,
        pady=5,
        relief='raised',
        bd=2
    )
    btn_cancel.pack(side='left', padx=5)

    btn_back = tk.Button(
        btn_frame,
        text="العودة",
        command=lambda: [pwd_window.destroy(), root_window.deiconify()],
        font=arabic_font,
        bg='orange',
        fg='white',
        padx=25,
        pady=5,
        relief='raised',
        bd=2
    )
    btn_back.pack(side='left', padx=5)


# Test function if run directly
if __name__ == "__main__":
    # Create a test window to demonstrate the functionality
    test_root = tk.Tk()
    test_root.withdraw()  # Hide the main window

    # Define a sample Arabic font
    arabic_font_test = ("Arial", 18)

    # Test the password prompt
    prompt_for_password(test_root, arabic_font_test)

    test_root.mainloop()
